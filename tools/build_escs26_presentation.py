#!/usr/bin/env python3
"""Build a polished ESCS'26 live presentation deck."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageEnhance, ImageFilter, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports" / "presentation"
ASSET_DIR = OUT_DIR / "assets"
OUT_PATH = OUT_DIR / "escs26_pcb_defect_detection_presentation_polished.pptx"

FIG_DIR = ROOT / "reports" / "publication" / "camera_ready_figures"
EXAMPLE_DIR = ROOT / "online_deployment" / "examples"

SLIDE_W = 13.333
SLIDE_H = 7.5

NAVY = RGBColor(16, 31, 45)
INK = RGBColor(24, 38, 46)
MUTED = RGBColor(83, 101, 111)
GREEN = RGBColor(0, 96, 76)
TEAL = RGBColor(0, 173, 166)
BLUE = RGBColor(32, 95, 175)
ORANGE = RGBColor(235, 91, 0)
GOLD = RGBColor(230, 165, 0)
BG = RGBColor(245, 248, 247)
CARD = RGBColor(255, 255, 255)
LINE = RGBColor(207, 222, 218)
WHITE = RGBColor(255, 255, 255)
PALE_GREEN = RGBColor(229, 246, 241)
PALE_ORANGE = RGBColor(255, 240, 226)
PALE_BLUE = RGBColor(232, 240, 252)


def rgb_tuple(color: RGBColor) -> tuple[int, int, int]:
    return color[0], color[1], color[2]


def font(size: int = 18, bold: bool = False, color: RGBColor = INK, name: str = "Aptos"):
    return {"size": Pt(size), "bold": bold, "color": color, "name": name}


def safe_font(size=36, bold=True):
    try:
        return ImageFont.truetype("Arial Bold.ttf" if bold else "Arial.ttf", size)
    except OSError:
        try:
            return ImageFont.truetype("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf", size)
        except OSError:
            return ImageFont.load_default()


def wrap_line_count(text: str, size_pt: float, max_width_in: float, bold: bool = True) -> int:
    """Estimate how many lines `text` wraps to at a given point size and box width.

    Uses real font metrics (via PIL) as a stand-in for the Aptos body font so the
    generator can size text containers instead of guessing a fixed height.
    """
    face = safe_font(size=max(8, round(size_pt)), bold=bold)
    max_width_pt = max_width_in * 72
    lines = 1
    current = ""
    for word in text.split():
        candidate = f"{current} {word}".strip()
        if current and face.getlength(candidate) > max_width_pt:
            lines += 1
            current = word
        else:
            current = candidate
    return lines


def set_bullet_glyph(paragraph, color: RGBColor, char: str = "▪", size_pct: int = 70, indent_in: float = 0.23):
    """Use native PowerPoint bullet formatting so bullets stay attached to their
    paragraph regardless of how many lines it wraps to."""
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buSzPct", "a:buFont"):
        existing = pPr.find(qn(tag))
        if existing is not None:
            pPr.remove(existing)
    bu_clr = OxmlElement("a:buClr")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", str(color))
    bu_clr.append(srgb)
    bu_sz = OxmlElement("a:buSzPct")
    bu_sz.set("val", str(size_pct * 1000))
    bu_font = OxmlElement("a:buFont")
    bu_font.set("typeface", "Arial")
    bu_char = OxmlElement("a:buChar")
    bu_char.set("char", char)
    for el in (bu_clr, bu_sz, bu_font, bu_char):
        pPr.append(el)
    indent = Inches(indent_in)
    pPr.set("marL", str(int(indent)))
    pPr.set("indent", str(int(-indent)))


def set_text(paragraph, text: str, *, size: int = 18, bold: bool = False, color: RGBColor = INK, name: str = "Aptos", align=None):
    paragraph.text = text
    paragraph.font.name = name
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align


def add_rect(slide, x, y, w, h, fill: RGBColor, line: RGBColor | None = None, radius=False, transparency: float = 0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def add_textbox(slide, text: str, x, y, w, h, *, size=18, bold=False, color=INK, name="Aptos", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    set_text(tf.paragraphs[0], text, size=size, bold=bold, color=color, name=name, align=align)
    return box


def add_background(slide, fill: RGBColor = BG):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill)


def add_header(slide, title: str, section: str, number: int):
    add_textbox(slide, section.upper(), 0.62, 0.25, 4.0, 0.22, size=8.5, bold=True, color=TEAL)
    add_textbox(slide, title, 0.62, 0.47, 9.9, 0.46, size=24, bold=True, color=NAVY, name="Aptos Display")
    add_rect(slide, 0.62, 1.08, 11.95, 0.025, TEAL)
    add_textbox(slide, f"ESCS'26 | {number}", 11.35, 7.05, 1.25, 0.22, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, fill=CARD, line=LINE):
    return add_rect(slide, x, y, w, h, fill, line=line, radius=True)


def add_chip(slide, text: str, x, y, w, fill=PALE_GREEN, color=GREEN):
    add_rect(slide, x, y, w, 0.36, fill, line=None, radius=True)
    add_textbox(slide, text, x + 0.12, y + 0.07, w - 0.24, 0.18, size=10.5, bold=True, color=color, align=PP_ALIGN.CENTER)


def add_metric(slide, label: str, value: str, note: str, x, y, w=2.55, fill=WHITE, accent=TEAL):
    add_card(slide, x, y, w, 1.2, fill=fill)
    add_rect(slide, x, y, 0.08, 1.2, accent, radius=False)
    add_textbox(slide, label.upper(), x + 0.24, y + 0.16, w - 0.34, 0.18, size=8.5, bold=True, color=MUTED)
    add_textbox(slide, value, x + 0.24, y + 0.38, w - 0.34, 0.35, size=24, bold=True, color=NAVY, name="Aptos Display")
    add_textbox(slide, note, x + 0.24, y + 0.82, w - 0.34, 0.2, size=9, color=MUTED)


def add_bullets(slide, bullets: list[str], x, y, w, h, *, size=17, color=INK, marker=TEAL):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, text in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(10)
        p.line_spacing = 1.08
        set_bullet_glyph(p, marker)
    return box


def add_table(slide, rows: list[list[str]], x, y, w, h, *, font_size=10.5, header=GREEN):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header
            else:
                cell.fill.fore_color.rgb = RGBColor(250, 252, 251) if r % 2 else WHITE
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(font_size)
                p.font.bold = r == 0
                p.font.color.rgb = WHITE if r == 0 else INK
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
    return table_shape


def add_big_quote(slide, text: str, x, y, w, h, fill=PALE_GREEN, accent=TEAL, size=18):
    inner_w = w - 0.55
    n_lines = wrap_line_count(text, size, inner_w, bold=True)
    line_height_in = size * 1.3 / 72
    needed_h = n_lines * line_height_in + 0.4
    h = max(h, needed_h)
    add_card(slide, x, y, w, h, fill=fill, line=accent)
    add_rect(slide, x, y, 0.11, h, accent)
    box = add_textbox(slide, text, x + 0.34, y + 0.2, inner_w, h - 0.35, size=size, bold=True, color=NAVY)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return h


def add_picture_fit(slide, path: Path, x, y, w, h):
    if not path.exists():
        add_big_quote(slide, f"Missing visual: {path.name}", x, y, w, h, fill=PALE_ORANGE, accent=ORANGE)
        return
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def trim_whitespace(input_path: Path, output_path: Path, padding=24):
    image = Image.open(input_path).convert("RGB")
    bg = Image.new("RGB", image.size, (255, 255, 255))
    diff = ImageChops.difference(image, bg)
    diff = ImageEnhance.Contrast(diff).enhance(4)
    bbox = diff.getbbox()
    if not bbox:
        image.save(output_path)
        return output_path
    left = max(0, bbox[0] - padding)
    top = max(0, bbox[1] - padding)
    right = min(image.width, bbox[2] + padding)
    bottom = min(image.height, bbox[3] + padding)
    image.crop((left, top, right, bottom)).save(output_path)
    return output_path


def make_hero_mosaic(output_path: Path) -> Path:
    examples = [
        ("Missing hole", EXAMPLE_DIR / "clean_missing_hole.jpg"),
        ("Mouse bite", EXAMPLE_DIR / "clean_mouse_bite.jpg"),
        ("Open circuit", EXAMPLE_DIR / "clean_open_circuit.jpg"),
        ("Short", EXAMPLE_DIR / "clean_short.jpg"),
        ("Spur", EXAMPLE_DIR / "clean_spur.jpg"),
        ("Spurious copper", EXAMPLE_DIR / "clean_spurious_copper.jpg"),
    ]
    canvas = Image.new("RGB", (1920, 1080), rgb_tuple(NAVY))
    draw = ImageDraw.Draw(canvas)
    tile_w, tile_h = 360, 250
    x0, y0 = 990, 120
    gap = 22
    label_font = safe_font(22, True)
    for idx, (label, path) in enumerate(examples):
        row, col = divmod(idx, 2)
        x = x0 + col * (tile_w + gap)
        y = y0 + row * (tile_h + gap)
        if path.exists():
            img = Image.open(path).convert("RGB")
            img = ImageOps.fit(img, (tile_w, tile_h), Image.Resampling.LANCZOS)
        else:
            img = Image.new("RGB", (tile_w, tile_h), (28, 55, 66))
        img = ImageEnhance.Color(img).enhance(1.25)
        img = ImageEnhance.Contrast(img).enhance(1.12)
        canvas.paste(img, (x, y))
        draw.rectangle((x, y + tile_h - 40, x + tile_w, y + tile_h), fill=(0, 0, 0))
        draw.text((x + 14, y + tile_h - 31), label, fill=(255, 255, 255), font=label_font)
    # Left-side subtle circuit lines and gradient panel.
    for i in range(13):
        y = 90 + i * 72
        draw.line((0, y, 890, y + 40), fill=(22, 74, 76), width=2)
    overlay = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    od.rectangle((0, 0, 960, 1080), fill=(4, 40, 35, 225))
    od.rectangle((890, 0, 1920, 1080), fill=(0, 0, 0, 45))
    canvas = Image.alpha_composite(canvas.convert("RGBA"), overlay).convert("RGB")
    canvas = canvas.filter(ImageFilter.UnsharpMask(radius=1.0, percent=120, threshold=3))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path)
    return output_path


def make_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "hero": make_hero_mosaic(ASSET_DIR / "pcb_hero_mosaic.png"),
        "success_trimmed": trim_whitespace(
            FIG_DIR / "class_balanced_success_examples.png",
            ASSET_DIR / "class_balanced_success_examples_trimmed.png",
            padding=10,
        ),
        "failure_trimmed": trim_whitespace(
            FIG_DIR / "representative_failure_cases.png",
            ASSET_DIR / "representative_failure_cases_trimmed.png",
            padding=10,
        ),
        "sample_pcb": EXAMPLE_DIR / "clean_open_circuit.jpg",
    }
    return assets


def content_slide(prs, number: int, title: str, section: str = "Study"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_header(slide, title, section, number)
    return slide


def add_section_label(slide, text: str, x, y, fill=PALE_GREEN, color=GREEN):
    add_rect(slide, x, y, 1.55, 0.34, fill, radius=True)
    add_textbox(slide, text.upper(), x + 0.12, y + 0.08, 1.31, 0.15, size=8, bold=True, color=color, align=PP_ALIGN.CENTER)


def slide_title(prs, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_picture_fit(slide, assets["hero"], 0, 0, SLIDE_W, SLIDE_H)
    add_section_label(slide, "ESCS'26 / CSCE'26", 0.7, 0.62, fill=RGBColor(221, 252, 244), color=GREEN)
    add_textbox(
        slide,
        "Real-Time PCB Defect Detection for Embedded Visual Inspection",
        0.72,
        1.5,
        5.6,
        2.05,
        size=32,
        bold=True,
        color=WHITE,
        name="Aptos Display",
    )
    add_textbox(
        slide,
        "A same-split comparison of YOLO11s and RT-DETR-L",
        0.78,
        3.75,
        6.4,
        0.35,
        size=18,
        color=RGBColor(215, 246, 240),
    )
    add_rect(slide, 0.78, 4.48, 1.15, 0.045, TEAL)
    add_textbox(
        slide,
        "Aditya Varun Dhayapulay and Paul Salvador Inventado\nCalifornia State University, Fullerton",
        0.78,
        5.65,
        6.5,
        0.7,
        size=14,
        bold=True,
        color=WHITE,
    )
    add_textbox(slide, "1", 12.15, 7.05, 0.55, 0.22, size=8, color=RGBColor(220, 235, 232), align=PP_ALIGN.RIGHT)


def slide_problem(prs, assets):
    slide = content_slide(prs, 2, "Why PCB Inspection Is Hard", "Motivation")
    add_picture_fit(slide, assets["sample_pcb"], 0.78, 1.42, 5.35, 4.65)
    add_rect(slide, 0.78, 5.62, 5.35, 0.45, RGBColor(0, 0, 0), transparency=30)
    add_textbox(slide, "Small visual defects must become actionable locations", 1.05, 5.76, 4.75, 0.18, size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_metric(slide, "Inspection risk", "small defects", "missing holes, shorts, spurs", 6.55, 1.45, 2.75, fill=WHITE, accent=ORANGE)
    add_metric(slide, "Output needed", "boxes + classes", "not image-level labels", 9.6, 1.45, 2.75, fill=WHITE, accent=TEAL)
    add_bullets(
        slide,
        [
            "Defects can become more expensive after assembly.",
            "Localization supports review, repair, and process analytics.",
            "A useful detector must balance recall, box quality, and latency.",
        ],
        6.8,
        3.25,
        5.2,
        2.1,
        size=18,
    )


def slide_goal(prs):
    slide = content_slide(prs, 3, "Research Question", "Goal")
    add_big_quote(
        slide,
        "Under one controlled PCB defect split, how do YOLO11s and RT-DETR-L compare for real-time visual inspection?",
        0.82,
        1.45,
        11.7,
        0.95,
    )
    labels = [("Input", "PCB image"), ("Models", "YOLO11s vs RT-DETR-L"), ("Outputs", "boxes, classes, metrics")]
    xs = [0.95, 4.75, 8.55]
    for idx, (label, body) in enumerate(labels):
        add_card(slide, xs[idx], 3.05, 3.05, 1.55, fill=WHITE)
        add_textbox(slide, label.upper(), xs[idx] + 0.22, 3.22, 2.6, 0.18, size=9, bold=True, color=TEAL)
        add_textbox(slide, body, xs[idx] + 0.22, 3.55, 2.6, 0.45, size=18, bold=True, color=NAVY, name="Aptos Display")
    add_bullets(
        slide,
        [
            "Deployment-oriented: accuracy, localization, latency, and failure behavior.",
            "Contribution is a reproducible benchmark, not a new detector architecture.",
        ],
        1.15,
        5.25,
        10.7,
        1.1,
        size=18,
    )


def slide_contributions(prs):
    slide = content_slide(prs, 4, "What the Camera-Ready Version Adds", "Contributions")
    items = [
        ("01", "Preserves practical comparison", "YOLO11s-1280 vs RT-DETR-L-640"),
        ("02", "Adds matched resolution", "YOLO11s-640 vs RT-DETR-L-640"),
        ("03", "Tests high-res RT-DETR-L", "1280 px, batch size 1"),
        ("04", "Fixes per-class semantics", "canonical class-ID mapping"),
        ("05", "Adds deployment artifact", "Hugging Face Space demo"),
    ]
    for idx, (num, title, body) in enumerate(items):
        x = 0.85 + (idx % 3) * 4.05
        y = 1.55 + (idx // 3) * 2.1
        add_card(slide, x, y, 3.55, 1.6, fill=WHITE)
        add_textbox(slide, num, x + 0.22, y + 0.18, 0.55, 0.28, size=16, bold=True, color=ORANGE, name="Aptos Display")
        add_textbox(slide, title, x + 0.88, y + 0.16, 2.5, 0.6, size=15, bold=True, color=NAVY)
        add_textbox(slide, body, x + 0.88, y + 0.78, 2.5, 0.4, size=11.5, color=MUTED)
    add_big_quote(slide, "Practical result + controlled evidence + honest deployment scope.", 2.0, 6.0, 9.2, 0.6, fill=PALE_BLUE, accent=BLUE)


def slide_dataset(prs):
    slide = content_slide(prs, 5, "Dataset and Defect Taxonomy", "Dataset")
    chips = ["Missing hole", "Mouse bite", "Open circuit", "Short", "Spur", "Spurious copper"]
    for idx, chip in enumerate(chips):
        add_chip(slide, chip, 0.82 + (idx % 3) * 2.1, 1.45 + (idx // 3) * 0.55, 1.75, fill=PALE_GREEN if idx % 2 == 0 else PALE_BLUE, color=GREEN if idx % 2 == 0 else BLUE)
    add_metric(slide, "Train images", "5,551", "deterministic split", 0.9, 3.05, 2.6, accent=TEAL)
    add_metric(slide, "Validation", "1,016", "2,106 instances", 3.8, 3.05, 2.6, accent=BLUE)
    add_metric(slide, "Test", "1,016", "2,179 instances", 6.7, 3.05, 2.6, accent=ORANGE)
    add_metric(slide, "Balancing", "800", "copy-paste train images", 9.6, 3.05, 2.6, accent=GOLD)
    add_big_quote(slide, "All reported configurations use the same validation and test image assignments.", 1.15, 5.25, 10.8, 0.75)


def slide_protocol(prs):
    slide = content_slide(prs, 6, "Reproducible Evaluation Protocol", "Protocol")
    steps = [
        ("Hardware", "Tesla V100-SXM2-32GB"),
        ("Software", "Ultralytics 8.4.51 / PyTorch 2.5.1"),
        ("Evaluation", "batch 1, workers 0, no TTA"),
        ("Metrics", "P, R, F1, mAP50, mAP50-95, latency"),
    ]
    for idx, (title, body) in enumerate(steps):
        x = 0.82 + idx * 3.08
        add_card(slide, x, 1.55, 2.65, 1.7)
        add_textbox(slide, f"{idx + 1}", x + 0.18, 1.73, 0.35, 0.28, size=16, bold=True, color=TEAL, name="Aptos Display")
        add_textbox(slide, title, x + 0.62, 1.73, 1.75, 0.24, size=13, bold=True, color=NAVY)
        add_textbox(slide, body, x + 0.24, 2.27, 2.1, 0.5, size=11.3, color=MUTED)
    add_big_quote(slide, "Baseline gate: archived checkpoint metrics must reproduce within 0.001 before new experiments are trusted.", 1.0, 4.25, 11.1, 0.85, fill=PALE_ORANGE, accent=ORANGE)
    add_textbox(slide, "Observed differences were approximately 10⁻⁸, confirming exact split/evaluator reproduction.", 1.2, 5.45, 10.6, 0.4, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def slide_models(prs):
    slide = content_slide(prs, 7, "Models and Training Settings", "Protocol")
    add_table(
        slide,
        [
            ["Configuration", "Image size", "Epochs", "Batch", "Purpose"],
            ["YOLO11s original", "1280", "50", "12", "accepted practical"],
            ["RT-DETR-L original", "640", "10", "4", "accepted practical"],
            ["YOLO11s controlled", "640", "50", "12", "matched resolution"],
            ["RT-DETR-L controlled", "640", "10", "4", "matched resolution"],
            ["RT-DETR-L high-res", "1280", "10", "1", "high-res analysis"],
        ],
        0.6,
        1.45,
        12.15,
        3.55,
        font_size=11,
    )
    add_big_quote(slide, "The 1280-pixel RT-DETR-L run completed; the planned 960-pixel fallback was not needed.", 1.2, 5.55, 10.8, 0.75)


def slide_resolution(prs):
    slide = content_slide(prs, 8, "The Key Fairness Issue: Resolution", "Reviewer response")
    add_card(slide, 0.8, 1.45, 5.35, 3.7, fill=WHITE)
    add_textbox(slide, "Accepted practical setup", 1.1, 1.78, 4.6, 0.34, size=18, bold=True, color=NAVY, name="Aptos Display")
    add_metric(slide, "YOLO11s", "1280 px", "more spatial detail", 1.15, 2.35, 2.15, accent=TEAL)
    add_metric(slide, "RT-DETR-L", "640 px", "lower input resolution", 3.65, 2.35, 2.15, accent=ORANGE)
    add_textbox(slide, "Useful as a deployment comparison, but not an architecture-only ablation.", 1.15, 4.0, 4.45, 0.55, size=14, color=MUTED)
    add_card(slide, 7.1, 1.45, 5.35, 3.7, fill=WHITE)
    add_textbox(slide, "Camera-ready fix", 7.4, 1.78, 4.6, 0.34, size=18, bold=True, color=NAVY, name="Aptos Display")
    add_metric(slide, "Matched", "640 vs 640", "same split and evaluator", 7.45, 2.35, 2.2, accent=BLUE)
    add_metric(slide, "High-res RT", "1280 px", "separate analysis", 9.95, 2.35, 2.2, accent=GOLD)
    add_textbox(slide, "Separates practical configurations from resolution-controlled evidence.", 7.45, 4.0, 4.45, 0.55, size=14, color=MUTED)


def add_result_bars(slide, x, y, title, yolo_value, rt_value, metric_label, max_value=1.0, yolo_label="YOLO11s", rt_label="RT-DETR-L"):
    add_textbox(slide, title, x, y, 3.6, 0.25, size=13, bold=True, color=NAVY)
    bar_w = 3.2
    y1 = y + 0.45
    for label, value, color, yy in [(yolo_label, yolo_value, TEAL, y1), (rt_label, rt_value, ORANGE, y1 + 0.48)]:
        add_textbox(slide, label, x, yy - 0.02, 1.15, 0.18, size=8.5, bold=True, color=MUTED)
        add_rect(slide, x + 1.25, yy, bar_w, 0.18, RGBColor(224, 233, 230), radius=True)
        add_rect(slide, x + 1.25, yy, bar_w * min(value / max_value, 1.0), 0.18, color, radius=True)
        add_textbox(slide, f"{value:.3f}" if value < 10 else f"{value:.1f}", x + 4.6, yy - 0.04, 0.6, 0.2, size=9.5, bold=True, color=NAVY)
    add_textbox(slide, metric_label, x, y + 1.18, 4.7, 0.2, size=8.5, color=MUTED)


def slide_practical(prs):
    slide = content_slide(prs, 9, "Accepted Practical Comparison", "Results")
    add_table(
        slide,
        [
            ["Model", "P", "R", "F1", "mAP50", "mAP50-95", "Total ms"],
            ["YOLO11s 1280", "0.880", "0.866", "0.873", "0.902", "0.502", "15.2"],
            ["RT-DETR-L 640", "0.886", "0.839", "0.862", "0.887", "0.470", "49.8"],
        ],
        0.68,
        1.35,
        11.95,
        1.45,
        font_size=12,
    )
    add_result_bars(slide, 0.95, 3.35, "Recall", 0.866, 0.839, "higher is better")
    add_result_bars(slide, 6.85, 3.35, "mAP50-95", 0.502, 0.470, "stricter localization")
    add_result_bars(slide, 0.95, 5.35, "Total latency", 15.2, 49.8, "milliseconds; lower is better", max_value=55, yolo_label="YOLO11s", rt_label="RT-DETR-L")
    add_big_quote(slide, "Interpretation: practical deployment choice among completed configurations.", 6.85, 5.43, 5.2, 0.68)


def slide_matched(prs):
    slide = content_slide(prs, 10, "Matched-Resolution Results", "Results")
    add_table(
        slide,
        [
            ["Model", "P", "R", "F1", "mAP50", "mAP50-95", "Total ms"],
            ["YOLO11s 640", "0.870", "0.852", "0.861", "0.892", "0.482", "13.2"],
            ["RT-DETR-L 640", "0.867", "0.839", "0.853", "0.881", "0.462", "49.0"],
        ],
        0.68,
        1.35,
        11.95,
        1.45,
        font_size=12,
    )
    add_result_bars(slide, 0.95, 3.35, "Recall", 0.852, 0.839, "same 640 px resolution")
    add_result_bars(slide, 6.85, 3.35, "mAP50-95", 0.482, 0.462, "same evaluator")
    add_big_quote(slide, "YOLO11s remains slightly stronger on recall and strict mAP while remaining much faster in this measured setup.", 1.15, 5.55, 10.95, 0.72)


def slide_highres(prs):
    slide = content_slide(prs, 11, "High-Resolution RT-DETR-L Finding", "Results")
    add_table(
        slide,
        [
            ["Model", "P", "R", "mAP50", "mAP50-95", "Total ms"],
            ["RT-DETR-L 640", "0.867", "0.839", "0.881", "0.462", "49.0"],
            ["RT-DETR-L 1280", "0.766", "0.534", "0.604", "0.295", "74.4"],
        ],
        0.82,
        1.45,
        11.6,
        1.45,
        font_size=12,
    )
    add_metric(slide, "mAP50-95 changed", "0.462 -> 0.295", "did not improve", 1.0, 3.65, 3.4, fill=PALE_ORANGE, accent=ORANGE)
    add_metric(slide, "Total latency", "49.0 -> 74.4 ms", "higher cost", 4.95, 3.65, 3.4, fill=PALE_ORANGE, accent=ORANGE)
    add_metric(slide, "Interpretation", "separate row", "not the matched comparison", 8.9, 3.65, 3.4, fill=PALE_BLUE, accent=BLUE)
    add_textbox(slide, "More pixels alone did not compensate for the fixed high-resolution training setup and computational cost.", 1.1, 5.65, 11.0, 0.42, size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def slide_perclass(prs):
    slide = content_slide(prs, 12, "Per-Class Behavior at 640 px", "Results")
    rows = [
        ["Class", "YOLO R", "RT R", "YOLO mAP50-95", "RT mAP50-95"],
        ["Missing hole", "0.986", "1.000", "0.569", "0.567"],
        ["Mouse bite", "0.763", "0.765", "0.403", "0.403"],
        ["Open circuit", "0.850", "0.880", "0.525", "0.513"],
        ["Short", "0.892", "0.865", "0.551", "0.505"],
        ["Spur", "0.767", "0.729", "0.381", "0.342"],
        ["Spurious copper", "0.853", "0.795", "0.462", "0.443"],
    ]
    add_table(slide, rows, 0.75, 1.35, 7.2, 4.55, font_size=10.5)
    add_card(slide, 8.45, 1.55, 3.95, 1.3, fill=PALE_GREEN)
    add_textbox(slide, "Strong classes", 8.72, 1.8, 3.3, 0.25, size=16, bold=True, color=GREEN)
    add_textbox(slide, "Missing hole, Open circuit, Short", 8.72, 2.18, 3.15, 0.24, size=12, color=MUTED)
    add_card(slide, 8.45, 3.25, 3.95, 1.3, fill=PALE_ORANGE)
    add_textbox(slide, "Harder classes", 8.72, 3.5, 3.3, 0.25, size=16, bold=True, color=ORANGE)
    add_textbox(slide, "Mouse bite and Spur remain subtle", 8.72, 3.88, 3.15, 0.24, size=12, color=MUTED)
    add_big_quote(slide, "Per-class analysis matters because aggregate mAP can hide inspection risk.", 8.45, 5.18, 3.95, 0.72)


def slide_successes(prs, assets):
    slide = content_slide(prs, 13, "Representative Successes", "Qualitative audit")
    add_card(slide, 0.68, 1.3, 8.1, 5.55, fill=WHITE)
    add_picture_fit(slide, assets["success_trimmed"], 0.9, 1.55, 7.65, 5.05)
    add_metric(slide, "Selection", "class-balanced", "matched YOLO11s-640", 9.25, 1.5, 3.0, accent=TEAL)
    add_bullets(
        slide,
        [
            "Shows localization across the six-class taxonomy.",
            "Bounding boxes support review and repair workflows.",
            "Visual audit complements aggregate metrics.",
        ],
        9.35,
        3.15,
        3.0,
        2.1,
        size=16,
    )


def slide_failures(prs, assets):
    slide = content_slide(prs, 14, "Representative Failure Cases", "Qualitative audit")
    add_card(slide, 0.68, 1.3, 8.1, 5.55, fill=WHITE)
    add_picture_fit(slide, assets["failure_trimmed"], 0.9, 1.55, 7.65, 5.05)
    add_metric(slide, "Selection score", "5FN + FP", "miss-weighted audit", 9.25, 1.5, 3.0, fill=PALE_ORANGE, accent=ORANGE)
    add_bullets(
        slide,
        [
            "False negatives are prioritized because missed defects are costly.",
            "Small/subtle classes motivate high-resolution reinspection.",
            "mAP50-95 exposes loose localization that mAP50 may hide.",
        ],
        9.35,
        3.15,
        3.0,
        2.35,
        size=16,
        marker=ORANGE,
    )


def slide_deployment(prs, assets):
    slide = content_slide(prs, 15, "Online Deployment Artifact", "Deployment")
    add_card(slide, 0.82, 1.35, 5.7, 4.9, fill=NAVY, line=NAVY)
    add_picture_fit(slide, assets["sample_pcb"], 1.15, 1.75, 5.05, 3.55)
    add_textbox(slide, "Hugging Face Space", 1.2, 5.55, 2.8, 0.25, size=15, bold=True, color=WHITE)
    add_textbox(slide, "adiivd-pcb-defect-detection.hf.space", 1.2, 5.9, 4.5, 0.2, size=10.5, color=RGBColor(200, 232, 226))
    quote_h = add_big_quote(slide, "Browser workflow: upload PCB image, run detector, inspect annotations, download results.", 7.05, 1.45, 5.2, 0.9)
    bullets_y = 1.45 + quote_h + 0.3
    add_bullets(
        slide,
        [
            "Demonstrates packaging and usability.",
            "Reports class counts, confidence values, and annotated detections.",
            "Not an embedded-device benchmark.",
            "No Jetson/TensorRT latency claims.",
        ],
        7.25,
        bullets_y,
        4.8,
        7.2 - bullets_y,
        size=17,
    )


def slide_limitations(prs):
    slide = content_slide(prs, 16, "Limitations and Future Work", "Scope")
    add_card(slide, 0.85, 1.45, 5.65, 4.65, fill=WHITE)
    add_textbox(slide, "Current scope", 1.15, 1.78, 4.8, 0.3, size=18, bold=True, color=NAVY, name="Aptos Display")
    add_bullets(
        slide,
        [
            "V100 latency, not edge latency.",
            "Same split, but not fully resource-normalized.",
            "In-distribution data, not factory-domain validation.",
        ],
        1.2,
        2.45,
        4.8,
        1.8,
        size=16,
    )
    add_card(slide, 6.85, 1.45, 5.65, 4.65, fill=PALE_GREEN)
    add_textbox(slide, "Next steps", 7.15, 1.78, 4.8, 0.3, size=18, bold=True, color=GREEN, name="Aptos Display")
    add_bullets(
        slide,
        [
            "Real factory captures and lighting variation.",
            "Resource-normalized training and latency budgets.",
            "High-resolution crop reinspection.",
            "Jetson/TensorRT benchmarking when hardware is available.",
        ],
        7.2,
        2.45,
        4.8,
        2.25,
        size=16,
    )


def slide_takeaways(prs):
    slide = content_slide(prs, 17, "Takeaways", "Close")
    add_metric(slide, "Practical result", "YOLO11s-1280", "higher recall, mAP50-95, lower latency", 0.9, 1.55, 3.65, fill=WHITE, accent=TEAL)
    add_metric(slide, "Controlled result", "YOLO11s-640", "competitive at matched resolution", 4.85, 1.55, 3.65, fill=WHITE, accent=BLUE)
    add_metric(slide, "Deployment lesson", "compare carefully", "resolution and latency both matter", 8.8, 1.55, 3.65, fill=WHITE, accent=ORANGE)
    add_big_quote(slide, "A deployment-oriented benchmark should separate completed configurations from controlled ablations.", 1.1, 3.75, 11.1, 0.85)
    add_textbox(slide, "Thank you. I welcome your questions.", 1.1, 5.55, 11.1, 0.55, size=28, bold=True, color=NAVY, name="Aptos Display", align=PP_ALIGN.CENTER)
    add_textbox(slide, "Aditya Varun Dhayapulay | California State University, Fullerton", 1.1, 6.25, 11.1, 0.25, size=13, color=MUTED, align=PP_ALIGN.CENTER)


def build_deck() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = make_assets()
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide_title(prs, assets)
    slide_problem(prs, assets)
    slide_goal(prs)
    slide_contributions(prs)
    slide_dataset(prs)
    slide_protocol(prs)
    slide_models(prs)
    slide_resolution(prs)
    slide_practical(prs)
    slide_matched(prs)
    slide_highres(prs)
    slide_perclass(prs)
    slide_successes(prs, assets)
    slide_failures(prs, assets)
    slide_deployment(prs, assets)
    slide_limitations(prs)
    slide_takeaways(prs)

    prs.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_deck()
