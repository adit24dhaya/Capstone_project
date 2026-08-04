#!/usr/bin/env python3
"""Build Q&A backup slides: paper material that is not in the main ESCS'26 deck.

Each slide answers one likely audience question with numbers taken from the
camera-ready paper, the accepted submission draft, and the archived evaluation
CSVs, so they can be screen-shared instantly during Q&A.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_escs26_presentation_final import (  # noqa: E402
    BLUE,
    GREEN,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    PALE_BLUE,
    PALE_GRAY,
    PALE_GREEN,
    PALE_ORANGE,
    ROOT,
    SLIDE_H,
    SLIDE_W,
    TEAL,
    WHITE,
    add_background,
    add_bullets,
    add_card,
    add_label,
    add_rect,
    add_table,
    add_takeaway,
    add_textbox,
    content_slide,
    emphasize_table_cells,
)
from pptx import Presentation  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches  # noqa: E402

OUT_PATH = ROOT / "escs26_qa_backup_slides.pptx"


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    add_label(slide, "ESCS'26 / CSCE'26", 0.72, 0.62, 2.4, fill=WHITE, color=GREEN, size=15)
    add_textbox(slide, "Q&A Backup Slides", 0.72, 1.5, 11.8, 0.9, size=44, bold=True, color=WHITE)
    add_textbox(
        slide,
        "Paper figures, tables, and analyses not shown in the main deck",
        0.75,
        2.55,
        11.5,
        0.5,
        size=22,
        color=PALE_GREEN,
    )
    add_rect(slide, 0.75, 3.3, 1.35, 0.05, TEAL)
    add_bullets(
        slide,
        [
            "B1  What changed between the accepted draft and the camera-ready",
            "B2  Why per-class numbers were replaced (label-mapping correction)",
            "B3  Evaluation-only diagnostic: 1280 checkpoint scored at 640",
            "B4  Validation-split results",
            "B5  Full matched-640 per-class table (all metrics, both models)",
            "B6  Fusion and inspection-cost supporting analysis",
            "B7  Latency: what is measured and how the numbers evolved",
            "B8  Dataset provenance and balancing  |  B9  Environment and reproducibility",
        ],
        0.78,
        3.75,
        11.7,
        3.3,
        size=18,
        color=WHITE,
        marker=TEAL,
        gap=6,
    )


def slide_versions(prs):
    slide = content_slide(prs, "B1", "What Changed From the Accepted Draft", "Backup / Versions")
    add_card(slide, 0.78, 1.35, 5.9, 4.6, fill=WHITE)
    add_textbox(slide, "ADDED FOR CAMERA-READY", 1.08, 1.66, 5.2, 0.27, size=16, bold=True, color=TEAL)
    add_bullets(
        slide,
        [
            "Matched-resolution YOLO11s-640 vs RT-DETR-L-640 experiment",
            "Completed RT-DETR-L 1280 px run (batch 1; no 960 px fallback needed)",
            "Corrected canonical class-ID mapping for per-class analysis",
            "Class-balanced success and failure example figures",
            "Complete manifest: counts, versions, seeds, checkpoints, tables, figures",
        ],
        1.1,
        2.14,
        5.35,
        3.6,
        size=18,
        gap=9,
    )
    add_card(slide, 6.95, 1.35, 5.62, 4.6, fill=PALE_BLUE, line=BLUE)
    add_textbox(slide, "PRESERVED FROM THE ACCEPTED STUDY", 7.25, 1.66, 5.0, 0.27, size=16, bold=True, color=BLUE)
    add_bullets(
        slide,
        [
            "Accepted practical comparison: YOLO11s-1280 vs RT-DETR-L-640",
            "Aggregate baseline metrics unchanged (reproduced within 0.001)",
            "WBF fusion and inspection-cost analysis kept as supporting context",
        ],
        7.27,
        2.14,
        5.1,
        2.4,
        size=18,
        marker=BLUE,
        gap=9,
    )
    add_textbox(
        slide,
        "Both reviewers asked for resolution control; the additions answer that directly.",
        7.27,
        4.9,
        5.1,
        0.8,
        size=17,
        bold=True,
        color=NAVY,
    )
    add_takeaway(slide, "Camera-ready = accepted baseline preserved + controlled evidence added.", y=6.18)


def slide_label_fix(prs):
    slide = content_slide(prs, "B2", "Why Per-Class Numbers Were Replaced", "Backup / Integrity")
    add_card(slide, 0.78, 1.35, 11.8, 1.62, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "WHAT WE FOUND DURING CAMERA-READY REPRODUCIBILITY WORK", 1.08, 1.58, 11.2, 0.26, size=16, bold=True, color=ORANGE)
    add_textbox(
        slide,
        "An earlier dataset export changed the displayed class-name order without remapping label IDs. Per-class rows for "
        "Missing_hole, Mouse_bite, and Spur could be misnamed; class-agnostic aggregate metrics were unaffected.",
        1.08,
        1.95,
        11.2,
        0.9,
        size=19,
        bold=True,
        color=NAVY,
    )
    add_card(slide, 0.78, 3.25, 5.85, 2.55, fill=WHITE)
    add_textbox(slide, "ACCEPTED-DRAFT PER-CLASS TABLE", 1.08, 3.52, 5.2, 0.26, size=16, bold=True, color=MUTED)
    add_bullets(
        slide,
        [
            "YOLO11s-1280 rows under the legacy display order",
            "Row names for three classes are not trustworthy",
            "Withdrawn rather than reinterpreted",
        ],
        1.1,
        3.98,
        5.3,
        1.7,
        size=18,
        marker=ORANGE,
        gap=8,
    )
    add_card(slide, 6.9, 3.25, 5.68, 2.55, fill=PALE_GREEN, line=TEAL)
    add_textbox(slide, "CAMERA-READY PER-CLASS TABLE", 7.2, 3.52, 5.0, 0.26, size=16, bold=True, color=GREEN)
    add_bullets(
        slide,
        [
            "Canonical class-ID mapping, validated before training",
            "Matched 640 px models trained on corrected labels",
            "Two manifests: legacy + canonical",
        ],
        7.22,
        3.98,
        5.15,
        1.7,
        size=18,
        gap=8,
    )
    add_takeaway(
        slide,
        "Per-class claims come only from corrected canonical labels; accepted aggregates remain valid and archived.",
        y=6.12,
        size=17,
    )


def slide_resize_diagnostic(prs):
    slide = content_slide(prs, "B3", "Can You Just Score the 1280 Checkpoint at 640?", "Backup / Diagnostic")
    table = add_table(
        slide,
        [
            ["YOLO11s-1280 checkpoint", "Recall", "mAP50-95"],
            ["Evaluated at 1280 px (as trained)", "0.866", "0.502"],
            ["Evaluated at 640 px (no retraining)", "0.810", "0.453"],
        ],
        1.3,
        1.5,
        10.7,
        1.62,
        widths=[6.1, 2.3, 2.3],
        font_size=18,
    )
    emphasize_table_cells(table, [(1, 1), (1, 2)], color=GREEN)
    add_card(slide, 1.3, 3.5, 5.15, 2.1, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "WHY THIS IS NOT THE MATCHED COMPARISON", 1.6, 3.78, 4.6, 0.5, size=16, bold=True, color=ORANGE)
    add_textbox(
        slide,
        "Resizing the input at evaluation time changes what the model sees without letting it adapt.",
        1.6,
        4.4,
        4.6,
        0.95,
        size=19,
        bold=True,
        color=NAVY,
    )
    add_card(slide, 6.85, 3.5, 5.15, 2.1, fill=PALE_GREEN, line=TEAL)
    add_textbox(slide, "WHAT WE DID INSTEAD", 7.15, 3.78, 4.6, 0.5, size=16, bold=True, color=GREEN)
    add_textbox(
        slide,
        "Trained YOLO11s from scratch at 640 px so the matched comparison uses resolution-specific training.",
        7.15,
        4.4,
        4.6,
        0.95,
        size=19,
        bold=True,
        color=NAVY,
    )
    add_takeaway(slide, "Resize-only evaluation is not a substitute for resolution-specific training.", y=6.1)


def slide_validation(prs):
    slide = content_slide(prs, "B4", "Validation-Split Results", "Backup / Results")
    add_label(slide, "ARCHIVED-CHECKPOINT EVALUATION  |  BATCH 1  |  SAME EVALUATOR", 3.0, 1.16, 7.3, fill=PALE_BLUE, color=BLUE, size=15)
    table = add_table(
        slide,
        [
            ["Model", "Split", "Precision", "Recall", "mAP50", "mAP50-95"],
            ["YOLO11s 1280", "Validation", "0.874", "0.874", "0.908", "0.505"],
            ["YOLO11s 1280", "Test", "0.880", "0.866", "0.902", "0.502"],
            ["RT-DETR-L 640", "Validation", "0.884", "0.855", "0.905", "0.472"],
            ["RT-DETR-L 640", "Test", "0.886", "0.839", "0.887", "0.470"],
        ],
        0.9,
        1.72,
        11.5,
        2.9,
        widths=[2.9, 2.0, 1.65, 1.65, 1.65, 1.65],
        font_size=17,
    )
    emphasize_table_cells(table, [(1, 3), (2, 3), (1, 5), (2, 5)], color=GREEN)
    add_bullets(
        slide,
        [
            "Validation split: 1,016 images, 2,106 instances. Test split: 1,016 images, 2,179 instances.",
            "The pattern matches on both splits: RT-DETR-L fractionally ahead on precision; YOLO11s ahead on recall and strict mAP.",
            "Test metrics were never used for hyperparameter tuning.",
        ],
        1.0,
        4.85,
        11.3,
        1.2,
        size=18,
        gap=7,
    )
    add_takeaway(slide, "Validation and test agree, so the headline test result is not a split artifact.", y=6.25, size=18)


def slide_per_class_full(prs):
    slide = content_slide(prs, "B5", "Full Matched-640 Per-Class Table", "Backup / Results")
    table = add_table(
        slide,
        [
            ["Class", "Model", "Precision", "Recall", "mAP50", "mAP50-95"],
            ["Missing_hole", "YOLO11s 640", "0.990", "0.986", "0.994", "0.569"],
            ["Missing_hole", "RT-DETR-L 640", "0.962", "1.000", "0.995", "0.567"],
            ["Mouse_bite", "YOLO11s 640", "0.821", "0.763", "0.835", "0.403"],
            ["Mouse_bite", "RT-DETR-L 640", "0.818", "0.765", "0.829", "0.403"],
            ["Open_circuit", "YOLO11s 640", "0.848", "0.850", "0.890", "0.525"],
            ["Open_circuit", "RT-DETR-L 640", "0.882", "0.880", "0.917", "0.513"],
            ["Short", "YOLO11s 640", "0.880", "0.892", "0.942", "0.551"],
            ["Short", "RT-DETR-L 640", "0.881", "0.865", "0.900", "0.505"],
            ["Spur", "YOLO11s 640", "0.869", "0.767", "0.839", "0.381"],
            ["Spur", "RT-DETR-L 640", "0.818", "0.729", "0.794", "0.342"],
            ["Spurious_copper", "YOLO11s 640", "0.815", "0.853", "0.852", "0.462"],
            ["Spurious_copper", "RT-DETR-L 640", "0.844", "0.795", "0.852", "0.443"],
        ],
        0.72,
        1.28,
        9.0,
        5.35,
        widths=[2.15, 2.05, 1.2, 1.2, 1.2, 1.2],
        font_size=13,
    )
    emphasize_table_cells(table, [(3, 5), (4, 5), (9, 5), (10, 5)], color=ORANGE)
    add_card(slide, 9.95, 1.35, 2.62, 2.2, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "HARDEST", 10.22, 1.62, 2.1, 0.24, size=15, bold=True, color=ORANGE)
    add_textbox(slide, "Spur 0.381\nMouse bite 0.403", 10.22, 2.0, 2.15, 0.7, size=18, bold=True, color=NAVY)
    add_textbox(slide, "(YOLO11s-640 mAP50-95)", 10.22, 2.95, 2.15, 0.45, size=13, color=MUTED)
    add_card(slide, 9.95, 3.8, 2.62, 2.2, fill=PALE_GREEN, line=TEAL)
    add_textbox(slide, "NOTE", 10.22, 4.07, 2.1, 0.24, size=15, bold=True, color=GREEN)
    add_textbox(
        slide,
        "Main deck shows recall + strict mAP only; this is the complete paper table.",
        10.22,
        4.45,
        2.15,
        1.4,
        size=15,
        color=INK,
    )
    add_textbox(
        slide,
        "Both detectors agree on which classes are hard, which points at data, not architecture.",
        0.9,
        6.75,
        10.3,
        0.32,
        size=17,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER,
    )


def slide_fusion(prs):
    slide = content_slide(prs, "B6", "Fusion and Inspection-Cost Analysis", "Backup / Supporting")
    add_label(slide, "ACCEPTED-STUDY SUPPORTING BUNDLE  |  NOT RE-RUN IN THE UNIFIED BATCH-1 EVALUATOR", 2.35, 1.14, 8.6, fill=PALE_GRAY, color=NAVY, size=14)
    table = add_table(
        slide,
        [
            ["Variant", "Precision", "Recall", "mAP50", "mAP50-95", "FP / image"],
            ["YOLO11s single model", "0.880", "0.866", "0.902", "0.502", "-"],
            ["Adaptive defect-aware WBF", "0.850", "0.903", "0.865", "0.493", "0.402"],
            ["Balanced WBF (permissive)", "0.757", "0.943", "0.903", "0.509", "0.852"],
        ],
        0.72,
        1.68,
        7.6,
        2.3,
        widths=[2.7, 1.05, 0.95, 0.95, 1.05, 0.9],
        font_size=14,
    )
    emphasize_table_cells(table, [(2, 2), (3, 2)], color=GREEN)
    add_card(slide, 8.6, 1.68, 3.98, 2.3, fill=WHITE)
    add_textbox(slide, "INSPECTION-COST SCORE", 8.88, 1.94, 3.5, 0.24, size=15, bold=True, color=TEAL)
    add_textbox(slide, "cost = FP + λ · FN  (per image)", 8.88, 2.32, 3.5, 0.35, size=18, bold=True, color=NAVY)
    add_textbox(
        slide,
        "λ = how much worse a missed defect is than a false alarm",
        8.88,
        2.88,
        3.5,
        0.9,
        size=15,
        color=MUTED,
    )
    table2 = add_table(
        slide,
        [
            ["λ", "Best variant", "Cost"],
            ["1", "YOLO11s single model", "0.672"],
            ["2", "Hybrid balanced WBF", "0.897"],
            ["5", "Hybrid balanced WBF", "1.505"],
            ["10", "Hybrid balanced WBF", "2.519"],
        ],
        0.72,
        4.25,
        6.4,
        2.0,
        widths=[0.8, 3.6, 2.0],
        font_size=14,
    )
    emphasize_table_cells(table2, [(1, 1)], color=GREEN)
    add_bullets(
        slide,
        [
            "Fusion buys recall (0.866 → 0.943) at the price of more false-positive review work.",
            "Single-model YOLO11s wins when misses and false alarms cost the same; ensembles win when misses dominate.",
        ],
        7.35,
        4.45,
        5.25,
        1.7,
        size=17,
        gap=8,
    )
    add_takeaway(slide, "Operating-point choice depends on the cost asymmetry of the inspection line.", y=6.5, size=17)


def slide_latency(prs):
    slide = content_slide(prs, "B7", "Latency: What Is Measured", "Backup / Protocol")
    add_textbox(
        slide,
        "Per image, batch 1, Tesla V100.  Total = preprocess + inference + postprocess.",
        0.9,
        1.2,
        11.5,
        0.32,
        size=18,
        bold=True,
        color=NAVY,
    )
    table = add_table(
        slide,
        [
            ["Configuration", "Inference (ms)", "Total (ms)"],
            ["YOLO11s 1280", "12.8", "15.2"],
            ["YOLO11s 640", "11.4", "13.2"],
            ["RT-DETR-L 640", "48.0 - 48.8", "49.0 - 49.8"],
            ["RT-DETR-L 1280", "72.4", "74.4"],
        ],
        0.9,
        1.75,
        7.2,
        2.75,
        widths=[3.2, 2.0, 2.0],
        font_size=16,
    )
    emphasize_table_cells(table, [(1, 2), (2, 2)], color=GREEN)
    add_card(slide, 8.5, 1.75, 4.08, 2.75, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "IF ASKED ABOUT THE ACCEPTED DRAFT", 8.78, 2.02, 3.6, 0.45, size=15, bold=True, color=ORANGE)
    add_textbox(
        slide,
        "The draft reported RT-DETR-L at 31.9 ms total from an earlier measurement pass. All camera-ready latencies were "
        "re-measured under the pinned batch-1, zero-worker protocol and supersede the draft numbers.",
        8.78,
        2.5,
        3.6,
        1.85,
        size=15,
        color=INK,
    )
    add_bullets(
        slide,
        [
            "The ranking never changed: YOLO11s is roughly 3x faster in every measurement pass.",
            "Latency numbers are V100 numbers; Jetson, TensorRT, RK3588, and FPGA remain unmeasured future work.",
        ],
        1.0,
        4.85,
        11.3,
        1.1,
        size=18,
        gap=8,
    )
    add_takeaway(slide, "Camera-ready latencies are the authoritative ones; the speed conclusion is stable across passes.", y=6.15, size=17)


def slide_dataset_provenance(prs):
    slide = content_slide(prs, "B8", "Dataset Provenance and Balancing", "Backup / Dataset")
    add_card(slide, 0.78, 1.35, 5.9, 4.6, fill=WHITE)
    add_textbox(slide, "PROVENANCE", 1.08, 1.66, 5.2, 0.27, size=16, bold=True, color=TEAL)
    add_bullets(
        slide,
        [
            "YOLO_PCB = project PCB dataset + overlapping DsPCBSD+ classes",
            "Accessed through the Kaggle dataset aditya2402/pcb-dataset",
            "Deterministic preparation: seed 42, stratified within source and class groups",
            "Two dataset manifests: legacy layout (aggregate reproduction) and canonical mapping (per-class claims)",
        ],
        1.1,
        2.14,
        5.35,
        3.5,
        size=18,
        gap=10,
    )
    add_card(slide, 6.95, 1.35, 5.62, 4.6, fill=PALE_GREEN, line=TEAL)
    add_textbox(slide, "COPY-PASTE BALANCING (TRAIN ONLY)", 7.25, 1.66, 5.0, 0.27, size=16, bold=True, color=GREEN)
    add_textbox(slide, "800 balancing images", 7.27, 2.12, 5.0, 0.45, size=28, bold=True, color=NAVY)
    add_bullets(
        slide,
        [
            "320 Missing_hole examples",
            "240 Mouse_bite examples",
            "240 Spur examples",
        ],
        7.27,
        2.75,
        5.1,
        1.5,
        size=19,
        gap=8,
    )
    add_textbox(
        slide,
        "Validation and test image assignments are identical for every reported configuration.",
        7.27,
        4.55,
        5.1,
        0.9,
        size=18,
        bold=True,
        color=NAVY,
    )
    add_takeaway(slide, "Balancing touches the training split only; evaluation data is never augmented.", y=6.18)


def slide_environment(prs):
    slide = content_slide(prs, "B9", "Environment and Reproducibility", "Backup / Protocol")
    add_table(
        slide,
        [
            ["Component", "Value"],
            ["GPU", "NVIDIA Tesla V100-SXM2-32GB"],
            ["Python", "3.12.8"],
            ["PyTorch", "2.5.1 + CUDA 12.4"],
            ["Ultralytics", "8.4.51"],
            ["OpenCV", "4.11.0"],
            ["NumPy", "1.26.4"],
        ],
        0.9,
        1.4,
        5.6,
        4.3,
        widths=[2.3, 3.3],
        font_size=16,
    )
    add_card(slide, 7.0, 1.4, 5.6, 4.3, fill=WHITE)
    add_textbox(slide, "REPRODUCIBILITY CONTROLS", 7.3, 1.7, 5.0, 0.27, size=16, bold=True, color=TEAL)
    add_bullets(
        slide,
        [
            "Baseline gate: archived checkpoints must reproduce within 0.001 before new runs (observed ~1e-8)",
            "Seed 42 for every controlled training run; one-epoch smoke test before each full configuration",
            "Checkpoint SHA-256 hashes recorded; all tables generated from saved evaluation CSVs",
            "Experiment manifest links counts, versions, hardware, seeds, checkpoints, tables, and figures",
        ],
        7.32,
        2.2,
        5.05,
        3.3,
        size=17,
        gap=9,
    )
    add_takeaway(slide, "Every number in the paper traces back to a saved CSV and a hashed checkpoint.", y=6.15)


BACKUP_SLIDES = (
    slide_cover,
    slide_versions,
    slide_label_fix,
    slide_resize_diagnostic,
    slide_validation,
    slide_per_class_full,
    slide_fusion,
    slide_latency,
    slide_dataset_provenance,
    slide_environment,
)


def add_backup_slides(prs):
    for slide_fn in BACKUP_SLIDES:
        slide_fn(prs)


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    add_backup_slides(prs)
    prs.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
