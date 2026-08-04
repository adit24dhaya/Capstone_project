#!/usr/bin/env python3
"""Build the final ESCS'26 Zoom presentation with an embedded offline demo."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageChops, ImageEnhance, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "reports" / "presentation" / "escs26_pcb_defect_detection_presentation_final.pptx"
SOURCE_DECKS = (
    ROOT / "reports" / "presentation" / "escs26_pcb_defect_detection_presentation.pptx",
    ROOT / "escs26_pcb_defect_detection_presentation_revised_with_demo.pptx",
)
ASSET_DIR = ROOT / "reports" / "presentation" / "final_assets"
FIG_DIR = ROOT / "reports" / "publication" / "camera_ready_figures"
EXAMPLE_DIR = ROOT / "online_deployment" / "examples"
VIDEO_PATH = ROOT / "reports" / "presentation" / "demo_video" / "demo_hf.mp4"
POSTER_PATH = ROOT / "reports" / "presentation" / "demo_video" / "demo_hf_poster.png"

SLIDE_W = 13.333
SLIDE_H = 7.5

NAVY = RGBColor(18, 43, 52)
INK = RGBColor(26, 43, 48)
MUTED = RGBColor(70, 91, 96)
GREEN = RGBColor(0, 92, 73)
TEAL = RGBColor(0, 162, 153)
BLUE = RGBColor(23, 104, 171)
ORANGE = RGBColor(222, 91, 0)
GOLD = RGBColor(218, 146, 0)
BG = RGBColor(246, 249, 248)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(202, 218, 214)
PALE_GREEN = RGBColor(229, 246, 241)
PALE_ORANGE = RGBColor(255, 239, 224)
PALE_BLUE = RGBColor(231, 241, 250)
PALE_GRAY = RGBColor(238, 243, 242)


def rgb_tuple(color: RGBColor) -> tuple[int, int, int]:
    return color[0], color[1], color[2]


def add_rect(slide, x, y, w, h, fill, line=None, rounded=False, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def add_textbox(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=20,
    bold=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.03,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.space_after = Pt(0)
    p.line_spacing = 1.0
    return box


def set_bullet(paragraph, color=TEAL, char="▪"):
    ppr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buSzPct", "a:buFont"):
        old = ppr.find(qn(tag))
        if old is not None:
            ppr.remove(old)
    bu_clr = OxmlElement("a:buClr")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", str(color))
    bu_clr.append(srgb)
    bu_size = OxmlElement("a:buSzPct")
    bu_size.set("val", "65000")
    bu_font = OxmlElement("a:buFont")
    bu_font.set("typeface", "Aptos")
    bu_char = OxmlElement("a:buChar")
    bu_char.set("char", char)
    for element in (bu_clr, bu_size, bu_font, bu_char):
        ppr.append(element)
    indent = Inches(0.26)
    ppr.set("marL", str(int(indent)))
    ppr.set("indent", str(int(-indent)))


def add_bullets(slide, bullets, x, y, w, h, *, size=20, color=INK, marker=TEAL, gap=7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for index, text in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = text
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.0
        set_bullet(p, marker)
    return box


def add_card(slide, x, y, w, h, fill=WHITE, line=LINE):
    return add_rect(slide, x, y, w, h, fill, line=line, rounded=True)


def add_background(slide, fill=BG):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill)


def add_header(slide, title, section, number):
    add_textbox(slide, section.upper(), 0.65, 0.22, 3.4, 0.23, size=13, bold=True, color=TEAL)
    add_textbox(slide, title, 0.65, 0.47, 11.4, 0.46, size=30, bold=True, color=NAVY)
    add_rect(slide, 0.65, 1.04, 12.0, 0.03, TEAL)
    add_textbox(slide, f"ESCS'26  |  {number}", 11.4, 7.1, 1.2, 0.18, size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def content_slide(prs, number, title, section):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_header(slide, title, section, number)
    return slide


def add_takeaway(slide, text, y=6.18, fill=PALE_GREEN, accent=TEAL, size=19):
    add_card(slide, 0.84, y, 11.65, 0.62, fill=fill, line=accent)
    add_rect(slide, 0.84, y, 0.1, 0.62, accent)
    add_textbox(
        slide,
        text,
        1.13,
        y + 0.1,
        11.02,
        0.4,
        size=size,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_label(slide, text, x, y, w, *, fill=PALE_GREEN, color=GREEN, size=14):
    add_rect(slide, x, y, w, 0.35, fill, rounded=True)
    add_textbox(slide, text, x + 0.08, y + 0.05, w - 0.16, 0.23, size=size, bold=True, color=color, align=PP_ALIGN.CENTER)


def add_stat(slide, label, value, note, x, y, w, *, accent=TEAL, fill=WHITE, value_size=25):
    add_card(slide, x, y, w, 1.12, fill=fill)
    add_rect(slide, x, y, 0.08, 1.12, accent)
    add_textbox(slide, label.upper(), x + 0.2, y + 0.12, w - 0.32, 0.2, size=13, bold=True, color=MUTED)
    add_textbox(slide, value, x + 0.2, y + 0.35, w - 0.32, 0.34, size=value_size, bold=True, color=NAVY)
    add_textbox(slide, note, x + 0.2, y + 0.78, w - 0.32, 0.2, size=14, color=MUTED)


def add_picture_contain(slide, path, x, y, w, h):
    path = Path(path)
    with Image.open(path) as image:
        ratio = image.width / image.height
    target_ratio = w / h
    if ratio > target_ratio:
        picture_w = w
        picture_h = w / ratio
        picture_x = x
        picture_y = y + (h - picture_h) / 2
    else:
        picture_h = h
        picture_w = h * ratio
        picture_x = x + (w - picture_w) / 2
        picture_y = y
    return slide.shapes.add_picture(
        str(path), Inches(picture_x), Inches(picture_y), width=Inches(picture_w), height=Inches(picture_h)
    )


def add_picture_cover(slide, path, x, y, w, h):
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    return picture


def add_table(slide, rows, x, y, w, h, *, widths=None, font_size=18, first_col_left=True):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    if widths:
        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)
    header_h = 0.52
    body_h = (h - header_h) / (len(rows) - 1)
    for row_index, row in enumerate(rows):
        table.rows[row_index].height = Inches(header_h if row_index == 0 else body_h)
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = GREEN
            else:
                cell.fill.fore_color.rgb = WHITE if row_index % 2 else PALE_GRAY
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(font_size)
                paragraph.font.bold = row_index == 0 or (col_index == 0 and row_index > 0)
                paragraph.font.color.rgb = WHITE if row_index == 0 else INK
                paragraph.alignment = PP_ALIGN.LEFT if first_col_left and col_index == 0 else PP_ALIGN.CENTER
                paragraph.space_after = Pt(0)
    return shape


def emphasize_table_cells(table_shape, coordinates, color=GREEN):
    table = table_shape.table
    for row, col in coordinates:
        for paragraph in table.cell(row, col).text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.color.rgb = color


def crop_panels(source, prefix):
    source = Path(source)
    image = Image.open(source).convert("RGB")
    # Crop the six image regions from the paper figure. These bounds retain every
    # prediction/ground-truth box while removing figure-level whitespace and tiny
    # filename captions that are unreadable in a Zoom window.
    if prefix == "success":
        boxes = [
            (145, 78, 575, 448),
            (895, 78, 1265, 448),
            (175, 505, 545, 878),
            (895, 505, 1265, 878),
            (175, 934, 545, 1310),
            (840, 934, 1315, 1310),
        ]
    else:
        boxes = [
            (175, 78, 545, 448),
            (895, 78, 1265, 448),
            (175, 440, 545, 875),
            (895, 440, 1265, 875),
            (175, 925, 545, 1312),
            (895, 925, 1265, 1312),
        ]
    paths = []
    for index, box in enumerate(boxes):
        out = ASSET_DIR / f"{prefix}_{index + 1}.png"
        image.crop(box).save(out)
        paths.append(out)
    return paths


def make_hero_mosaic(path):
    # One tile per defect class, each from a different board template so the
    # six tiles are visually distinct instead of six crops of the same board.
    dataset_images = ROOT / "Dataset" / "PCB-DATASET-master" / "images"
    examples = [
        ("Missing hole", dataset_images / "Missing_hole" / "01_missing_hole_01.jpg"),
        ("Mouse bite", dataset_images / "Mouse_bite" / "04_mouse_bite_01.jpg"),
        ("Open circuit", dataset_images / "Open_circuit" / "05_open_circuit_01.jpg"),
        ("Short", dataset_images / "Short" / "08_short_01.jpg"),
        ("Spur", dataset_images / "Spur" / "10_spur_01.jpg"),
        ("Spurious copper", dataset_images / "Spurious_copper" / "11_spurious_copper_01.jpg"),
    ]
    canvas = Image.new("RGB", (1920, 1080), rgb_tuple(NAVY))
    tile_w, tile_h = 350, 238
    left, top, gap = 1050, 135, 22
    for index, (_, source) in enumerate(examples):
        row, col = divmod(index, 2)
        image = Image.open(source).convert("RGB")
        image = ImageOps.fit(image, (tile_w, tile_h), Image.Resampling.LANCZOS)
        image = ImageEnhance.Contrast(image).enhance(1.08)
        canvas.paste(image, (left + col * (tile_w + gap), top + row * (tile_h + gap)))
    overlay = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    from PIL import ImageDraw

    draw = ImageDraw.Draw(overlay)
    draw.rectangle((0, 0, 1020, 1080), fill=(0, 62, 50, 245))
    draw.rectangle((1020, 0, 1920, 1080), fill=(0, 0, 0, 35))
    canvas = Image.alpha_composite(canvas.convert("RGBA"), overlay).convert("RGB")
    canvas.save(path)
    return path


def make_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    hero = make_hero_mosaic(ASSET_DIR / "title_pcb_mosaic.png")
    success = crop_panels(FIG_DIR / "class_balanced_success_examples.png", "success")
    failure = crop_panels(FIG_DIR / "representative_failure_cases.png", "failure")
    return {"hero": hero, "success": success, "failure": failure}


def slide_title(prs, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_picture_cover(slide, assets["hero"], 0, 0, SLIDE_W, SLIDE_H)
    add_label(slide, "ESCS'26 / CSCE'26", 0.72, 0.58, 2.4, fill=WHITE, color=GREEN, size=15)
    add_textbox(
        slide,
        "Real-Time PCB Defect Detection\nfor Embedded Visual Inspection",
        0.72,
        1.35,
        6.2,
        1.65,
        size=36,
        bold=True,
        color=WHITE,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide,
        "A Same-Split Comparison of YOLO11s and RT-DETR-L",
        0.75,
        3.34,
        6.15,
        0.56,
        size=21,
        bold=True,
        color=RGBColor(216, 245, 239),
    )
    add_rect(slide, 0.75, 4.18, 1.35, 0.05, TEAL)
    add_textbox(
        slide,
        "Aditya Varun Dhayapulay and Paul Salvador Inventado",
        0.75,
        5.25,
        6.45,
        0.35,
        size=18,
        bold=True,
        color=WHITE,
    )
    add_textbox(
        slide,
        "California State University, Fullerton",
        0.75,
        5.69,
        6.0,
        0.3,
        size=17,
        color=RGBColor(216, 245, 239),
    )


def slide_motivation(prs):
    slide = content_slide(prs, 2, "Why PCB Inspection Is Hard", "Motivation")
    add_card(slide, 0.78, 1.35, 5.25, 4.95, fill=NAVY, line=NAVY)
    add_picture_contain(slide, EXAMPLE_DIR / "clean_open_circuit.jpg", 0.98, 1.65, 4.85, 3.38)
    add_textbox(slide, "SMALL DEFECTS", 1.1, 5.27, 2.0, 0.25, size=16, bold=True, color=TEAL)
    add_textbox(slide, "Subtle visual changes must become actionable locations.", 1.1, 5.58, 4.55, 0.42, size=20, bold=True, color=WHITE)
    add_card(slide, 6.42, 1.35, 2.85, 1.42, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "INSPECTION RISK", 6.7, 1.61, 2.3, 0.23, size=16, bold=True, color=ORANGE)
    add_textbox(slide, "Missed defects cost more after assembly", 6.7, 1.94, 2.2, 0.55, size=20, bold=True, color=NAVY)
    add_card(slide, 9.57, 1.35, 2.85, 1.42, fill=PALE_GREEN, line=TEAL)
    add_textbox(slide, "OUTPUT NEEDED", 9.85, 1.61, 2.3, 0.23, size=16, bold=True, color=GREEN)
    add_textbox(slide, "Boxes and classes, not only image labels", 9.85, 1.94, 2.2, 0.55, size=20, bold=True, color=NAVY)
    add_bullets(
        slide,
        [
            "Classification alone cannot show where a defect is.",
            "Localization supports review, repair, and process analytics.",
            "Useful inspection balances recall, box quality, and latency.",
        ],
        6.65,
        3.28,
        5.4,
        2.35,
        size=21,
    )


def slide_question(prs):
    slide = content_slide(prs, 3, "Research Question", "Study")
    add_card(slide, 0.82, 1.35, 11.7, 1.18, fill=PALE_GREEN, line=TEAL)
    add_textbox(
        slide,
        "Under one controlled PCB defect split, how do YOLO11s and RT-DETR-L compare for real-time visual inspection?",
        1.12,
        1.58,
        11.1,
        0.72,
        size=27,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    flow = [("PCB image", TEAL), ("YOLO11s  |  RT-DETR-L", BLUE), ("Boxes, classes, metrics", ORANGE)]
    for index, (label, accent) in enumerate(flow):
        x = 0.95 + index * 4.15
        add_card(slide, x, 3.06, 3.45, 1.25, fill=WHITE, line=accent)
        add_textbox(slide, label, x + 0.2, 3.39, 3.05, 0.55, size=21, bold=True, color=NAVY, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if index < 2:
            add_textbox(slide, "→", x + 3.54, 3.38, 0.42, 0.4, size=28, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    add_label(slide, "Accuracy", 1.02, 4.9, 2.3, size=17)
    add_label(slide, "Localization", 3.62, 4.9, 2.3, fill=PALE_BLUE, color=BLUE, size=17)
    add_label(slide, "Latency", 6.22, 4.9, 2.3, fill=PALE_ORANGE, color=ORANGE, size=17)
    add_label(slide, "Failure behavior", 8.82, 4.9, 2.75, fill=PALE_GRAY, color=NAVY, size=17)
    add_takeaway(slide, "Reproducible deployment-oriented benchmark, not a new detector architecture.")


def slide_contributions(prs):
    slide = content_slide(prs, 4, "What This Study Contributes", "Contributions")
    items = [
        ("01", "Preserves practical comparison", "YOLO11s-1280 vs RT-DETR-L-640"),
        ("02", "Adds matched-resolution comparison", "YOLO11s-640 vs RT-DETR-L-640"),
        ("03", "Tests high-resolution RT-DETR-L", "1280 px with batch size 1"),
        ("04", "Fixes per-class semantic mapping", "Canonical class-ID interpretation"),
        ("05", "Adds an online deployment artifact", "Self-contained browser workflow"),
    ]
    positions = [(0.8, 1.35), (4.75, 1.35), (8.7, 1.35), (2.78, 3.62), (6.73, 3.62)]
    for (number, title, detail), (x, y) in zip(items, positions):
        add_card(slide, x, y, 3.62, 1.78, fill=WHITE)
        add_textbox(slide, number, x + 0.2, y + 0.19, 0.62, 0.4, size=25, bold=True, color=ORANGE)
        add_textbox(slide, title, x + 0.92, y + 0.17, 2.48, 0.62, size=18, bold=True, color=NAVY)
        add_textbox(slide, detail, x + 0.92, y + 0.96, 2.42, 0.48, size=16, color=MUTED)
    add_takeaway(slide, "Practical result + controlled evidence + honest deployment scope.", y=6.05)


def slide_dataset(prs):
    slide = content_slide(prs, 5, "Dataset and Defect Taxonomy", "Dataset")
    classes = [
        ("Missing hole", "clean_missing_hole.jpg"),
        ("Mouse bite", "clean_mouse_bite.jpg"),
        ("Open circuit", "clean_open_circuit.jpg"),
        ("Short", "clean_short.jpg"),
        ("Spur", "clean_spur.jpg"),
        ("Spurious copper", "clean_spurious_copper.jpg"),
    ]
    for index, (label, filename) in enumerate(classes):
        row, col = divmod(index, 3)
        x = 0.72 + col * 2.65
        y = 1.29 + row * 2.02
        add_card(slide, x, y, 2.42, 1.76, fill=WHITE)
        add_picture_contain(slide, EXAMPLE_DIR / filename, x + 0.09, y + 0.08, 2.24, 1.18)
        add_textbox(slide, label, x + 0.12, y + 1.39, 2.18, 0.24, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_card(slide, 8.75, 1.29, 3.82, 4.75, fill=WHITE)
    add_textbox(slide, "DETERMINISTIC SPLIT", 9.03, 1.56, 3.25, 0.24, size=16, bold=True, color=TEAL)
    stats = [
        ("5,551", "training images"),
        ("1,016", "validation images  |  2,106 instances"),
        ("1,016", "test images  |  2,179 instances"),
        ("800", "copy-paste balancing images"),
    ]
    for index, (value, label) in enumerate(stats):
        y = 2.0 + index * 0.92
        add_textbox(slide, value, 9.05, y, 1.1, 0.35, size=24, bold=True, color=NAVY)
        add_textbox(slide, label, 10.2, y + 0.04, 2.0, 0.42, size=16, color=MUTED)
    add_textbox(slide, "Balancing affects training only.", 9.05, 5.56, 3.08, 0.25, size=16, bold=True, color=ORANGE)
    add_takeaway(slide, "All reported configurations use the same validation and test image assignments.", y=6.2, size=18)


def slide_protocol(prs):
    slide = content_slide(prs, 6, "Reproducible Evaluation Protocol", "Protocol")
    cards = [
        ("1", "Hardware", "Tesla V100-SXM2-32GB"),
        ("2", "Software", "Ultralytics 8.4.51\nPyTorch 2.5.1"),
        ("3", "Evaluation", "Batch size 1\nWorkers 0  |  No TTA"),
        ("4", "Metrics", "Precision, Recall, F1\nmAP50, mAP50-95, latency"),
    ]
    for index, (number, title, detail) in enumerate(cards):
        x = 0.72 + index * 3.15
        add_card(slide, x, 1.38, 2.78, 2.05, fill=WHITE)
        add_textbox(slide, number, x + 0.2, 1.61, 0.45, 0.38, size=25, bold=True, color=TEAL)
        add_textbox(slide, title, x + 0.74, 1.64, 1.76, 0.3, size=20, bold=True, color=NAVY)
        add_textbox(slide, detail, x + 0.22, 2.26, 2.34, 0.75, size=17, color=MUTED, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_card(slide, 0.96, 4.02, 11.4, 1.45, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "BASELINE REPRODUCTION GATE", 1.27, 4.26, 5.0, 0.28, size=17, bold=True, color=ORANGE)
    add_textbox(slide, "Archived checkpoint metrics must reproduce within 0.001 before new experiments are trusted.", 1.27, 4.71, 10.65, 0.52, size=21, bold=True, color=NAVY)
    add_takeaway(slide, "Observed differences were approximately 10⁻⁸, confirming exact split and evaluator reproduction.", y=5.83, fill=PALE_BLUE, accent=BLUE, size=18)


def slide_models(prs):
    slide = content_slide(prs, 7, "Models and Training Settings", "Protocol")
    add_card(slide, 0.76, 1.24, 5.8, 0.76, fill=PALE_GREEN, line=TEAL)
    add_textbox(slide, "YOLO11s", 1.02, 1.45, 1.4, 0.3, size=18, bold=True, color=GREEN)
    add_textbox(slide, "One-stage CNN for low-latency inference", 2.35, 1.39, 3.85, 0.48, size=16, color=INK, valign=MSO_ANCHOR.MIDDLE)
    add_card(slide, 6.78, 1.24, 5.8, 0.76, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "RT-DETR-L", 7.04, 1.45, 1.75, 0.3, size=18, bold=True, color=ORANGE)
    add_textbox(slide, "Transformer detector with attention and object queries", 8.63, 1.34, 3.62, 0.58, size=16, color=INK, valign=MSO_ANCHOR.MIDDLE)
    table = add_table(
        slide,
        [
            ["Configuration", "Pixels", "Epochs", "Batch", "Purpose"],
            ["YOLO11s original", "1280", "50", "12", "Practical comparison"],
            ["RT-DETR-L original", "640", "10", "4", "Practical comparison"],
            ["YOLO11s controlled", "640", "50", "12", "Matched resolution"],
            ["RT-DETR-L controlled", "640", "10", "4", "Matched resolution"],
            ["RT-DETR-L high resolution", "1280", "10", "1", "High-resolution analysis"],
        ],
        0.7,
        2.2,
        11.95,
        3.56,
        widths=[3.2, 1.2, 1.25, 1.15, 5.15],
        font_size=17,
    )
    emphasize_table_cells(table, [(1, 0), (2, 0), (3, 0), (4, 0), (5, 0)], color=NAVY)
    add_takeaway(slide, "The 1280-pixel RT-DETR-L run completed; the planned 960-pixel fallback was not needed.", y=6.05, size=18)


def slide_fairness(prs):
    slide = content_slide(prs, 8, "Why Input Resolution Must Be Controlled", "Experiment design")
    add_card(slide, 0.72, 1.35, 5.85, 4.56, fill=WHITE)
    add_textbox(slide, "ACCEPTED PRACTICAL SETUP", 1.02, 1.66, 4.9, 0.25, size=17, bold=True, color=ORANGE)
    add_textbox(slide, "YOLO11s", 1.04, 2.28, 2.0, 0.3, size=21, bold=True, color=GREEN)
    add_textbox(slide, "1280 px", 3.42, 2.13, 2.48, 0.55, size=32, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "RT-DETR-L", 1.04, 3.15, 2.0, 0.3, size=21, bold=True, color=ORANGE)
    add_textbox(slide, "640 px", 3.42, 3.0, 2.48, 0.55, size=32, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_card(slide, 1.02, 4.16, 5.2, 1.2, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "Useful deployment comparison", 1.28, 4.4, 4.68, 0.28, size=20, bold=True, color=NAVY)
    add_textbox(slide, "Not an architecture-only ablation", 1.28, 4.78, 4.68, 0.28, size=18, color=ORANGE)
    add_card(slide, 6.84, 1.35, 5.77, 4.56, fill=WHITE)
    add_textbox(slide, "CAMERA-READY CONTROLLED EVIDENCE", 7.14, 1.66, 5.1, 0.25, size=17, bold=True, color=TEAL)
    add_textbox(slide, "640 vs 640", 7.24, 2.22, 4.98, 0.62, size=34, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Same split and evaluator", 7.24, 2.91, 4.98, 0.3, size=19, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_rect(slide, 7.5, 3.46, 4.45, 0.03, LINE)
    add_textbox(slide, "Separate high-resolution RT test", 7.24, 3.8, 4.98, 0.3, size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "RT-DETR-L at 1280 px", 7.24, 4.32, 4.98, 0.46, size=27, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "Separate practical configurations from resolution-controlled evidence.", y=6.12)


def add_result_table(slide, rows, winner_cells, *, y=1.35):
    table = add_table(
        slide,
        rows,
        0.67,
        y,
        11.98,
        1.52,
        widths=[2.55, 1.3, 1.3, 1.3, 1.45, 1.72, 1.78],
        font_size=18,
    )
    emphasize_table_cells(table, winner_cells, color=GREEN)
    return table


def comparison_tile(slide, title, left_label, left_value, right_label, right_value, x, *, note, lower=False):
    add_card(slide, x, 3.2, 3.72, 2.15, fill=WHITE)
    add_textbox(slide, title, x + 0.22, 3.43, 3.28, 0.28, size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, left_label, x + 0.2, 3.96, 1.5, 0.23, size=15, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_textbox(slide, right_label, x + 2.0, 3.96, 1.5, 0.23, size=15, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
    add_textbox(slide, left_value, x + 0.2, 4.28, 1.5, 0.48, size=28, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, right_value, x + 2.0, 4.28, 1.5, 0.48, size=28, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, note + ("  |  lower is better" if lower else "  |  higher is better"), x + 0.3, 4.92, 3.12, 0.22, size=14, color=MUTED, align=PP_ALIGN.CENTER)


def slide_practical(prs):
    slide = content_slide(prs, 9, "Practical Configuration Results", "Results")
    add_result_table(
        slide,
        [
            ["Model", "Precision", "Recall", "F1", "mAP50", "mAP50-95", "Total latency"],
            ["YOLO11s 1280", "0.880", "0.866", "0.873", "0.902", "0.502", "15.2 ms"],
            ["RT-DETR-L 640", "0.886", "0.839", "0.862", "0.887", "0.470", "49.8 ms"],
        ],
        [(1, 2), (1, 3), (1, 4), (1, 5), (1, 6), (2, 1)],
    )
    comparison_tile(slide, "Recall", "YOLO11s 1280", "0.866", "RT-DETR-L 640", "0.839", 0.72, note="detection coverage")
    comparison_tile(slide, "mAP50-95", "YOLO11s 1280", "0.502", "RT-DETR-L 640", "0.470", 4.81, note="strict localization")
    comparison_tile(slide, "Total latency", "YOLO11s 1280", "15.2", "RT-DETR-L 640", "49.8", 8.9, note="milliseconds", lower=True)
    add_takeaway(slide, "Among completed practical configurations, YOLO11s-1280 provided stronger recall, stricter localization, and lower measured latency.", y=5.84, size=17)


def slide_matched(prs):
    slide = content_slide(prs, 10, "Matched-Resolution Results", "Results")
    add_label(slide, "SAME 640-PIXEL RESOLUTION AND SAME EVALUATOR", 3.42, 1.14, 6.5, fill=PALE_BLUE, color=BLUE, size=16)
    add_result_table(
        slide,
        [
            ["Model", "Precision", "Recall", "F1", "mAP50", "mAP50-95", "Total latency"],
            ["YOLO11s 640", "0.870", "0.852", "0.861", "0.892", "0.482", "13.2 ms"],
            ["RT-DETR-L 640", "0.867", "0.839", "0.853", "0.881", "0.462", "49.0 ms"],
        ],
        [(1, 1), (1, 2), (1, 3), (1, 4), (1, 5), (1, 6)],
        y=1.57,
    )
    comparison_tile(slide, "Recall", "YOLO11s 640", "0.852", "RT-DETR-L 640", "0.839", 0.72, note="detection coverage")
    comparison_tile(slide, "mAP50-95", "YOLO11s 640", "0.482", "RT-DETR-L 640", "0.462", 4.81, note="strict localization")
    comparison_tile(slide, "Total latency", "YOLO11s 640", "13.2", "RT-DETR-L 640", "49.0", 8.9, note="milliseconds", lower=True)
    add_takeaway(slide, "At matched resolution, YOLO11s remained slightly stronger in recall and strict mAP while remaining much faster in this measured setup.", y=5.84, size=17)


def slide_highres(prs):
    slide = content_slide(prs, 11, "High-Resolution RT-DETR-L Finding", "Results")
    table = add_table(
        slide,
        [
            ["RT-DETR-L setting", "Precision", "Recall", "mAP50", "mAP50-95", "Total latency"],
            ["640 px", "0.867", "0.839", "0.881", "0.462", "49.0 ms"],
            ["1280 px", "0.766", "0.534", "0.604", "0.295", "74.4 ms"],
        ],
        0.78,
        1.34,
        11.78,
        1.58,
        widths=[2.8, 1.65, 1.65, 1.65, 1.85, 2.18],
        font_size=18,
    )
    emphasize_table_cells(table, [(1, 1), (1, 2), (1, 3), (1, 4), (1, 5)], color=GREEN)
    add_card(slide, 0.88, 3.32, 5.55, 2.16, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "STRICT LOCALIZATION", 1.2, 3.62, 4.9, 0.25, size=17, bold=True, color=ORANGE)
    add_textbox(slide, "0.462  →  0.295", 1.2, 4.02, 4.9, 0.55, size=34, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "mAP50-95 did not improve", 1.2, 4.73, 4.9, 0.3, size=19, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
    add_card(slide, 6.87, 3.32, 5.55, 2.16, fill=PALE_BLUE, line=BLUE)
    add_textbox(slide, "COMPUTATIONAL COST", 7.19, 3.62, 4.9, 0.25, size=17, bold=True, color=BLUE)
    add_textbox(slide, "49.0  →  74.4 ms", 7.19, 4.02, 4.9, 0.55, size=34, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Higher measured latency", 7.19, 4.73, 4.9, 0.3, size=19, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "More pixels alone did not compensate for this fixed high-resolution training setup; this is a separate analysis, not the matched comparison.", y=5.92, size=17)


def slide_per_class(prs):
    slide = content_slide(prs, 12, "Per-Class Behavior at 640 Pixels", "Results")
    table = add_table(
        slide,
        [
            ["Class", "YOLO Recall", "RT Recall", "YOLO mAP50-95", "RT mAP50-95"],
            ["Missing hole", "0.986", "1.000", "0.569", "0.567"],
            ["Mouse bite", "0.763", "0.765", "0.403", "0.403"],
            ["Open circuit", "0.850", "0.880", "0.525", "0.513"],
            ["Short", "0.892", "0.865", "0.551", "0.505"],
            ["Spur", "0.767", "0.729", "0.381", "0.342"],
            ["Spurious copper", "0.853", "0.795", "0.462", "0.443"],
        ],
        0.67,
        1.3,
        8.65,
        4.95,
        widths=[2.25, 1.45, 1.45, 1.75, 1.75],
        font_size=16,
    )
    emphasize_table_cells(table, [(1, 2), (2, 2), (3, 2), (4, 1), (5, 1), (6, 1), (1, 3), (3, 3), (4, 3), (5, 3), (6, 3)], color=GREEN)
    add_card(slide, 9.65, 1.38, 2.92, 1.68, fill=PALE_GREEN, line=TEAL)
    add_textbox(slide, "STRONG CLASSES", 9.92, 1.68, 2.38, 0.24, size=16, bold=True, color=GREEN)
    add_textbox(slide, "Missing hole\nOpen circuit\nShort", 9.92, 2.08, 2.35, 0.73, size=19, bold=True, color=NAVY)
    add_card(slide, 9.65, 3.34, 2.92, 1.68, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "HARDER CLASSES", 9.92, 3.64, 2.38, 0.24, size=16, bold=True, color=ORANGE)
    add_textbox(slide, "Mouse bite\nSpur", 9.92, 4.08, 2.35, 0.58, size=21, bold=True, color=NAVY)
    add_takeaway(slide, "Aggregate mAP can hide class-specific inspection risk.", y=6.38, size=18)


def add_panel_grid(slide, panels, labels):
    for index, (panel, label) in enumerate(zip(panels, labels)):
        row, col = divmod(index, 3)
        x = 0.68 + col * 2.84
        y = 1.27 + row * 2.45
        add_card(slide, x, y, 2.64, 2.24, fill=WHITE)
        add_picture_contain(slide, panel, x + 0.08, y + 0.08, 2.48, 1.76)
        add_textbox(slide, label, x + 0.08, y + 1.9, 2.48, 0.22, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def slide_successes(prs, assets):
    slide = content_slide(prs, 13, "Representative Successes", "Qualitative audit")
    add_panel_grid(
        slide,
        assets["success"],
        ["Missing hole", "Mouse bite + Open circuit", "Mixed copper defects", "Short + Spur", "Short + Spur", "Spurious copper"],
    )
    add_card(slide, 9.32, 1.34, 3.12, 4.75, fill=WHITE)
    add_textbox(slide, "CLASS-BALANCED SELECTION", 9.62, 1.68, 2.52, 0.4, size=16, bold=True, color=TEAL)
    add_textbox(slide, "Matched YOLO11s-640", 9.62, 2.28, 2.52, 0.6, size=23, bold=True, color=NAVY)
    add_bullets(
        slide,
        [
            "All six defect classes are represented.",
            "Boxes show localization and confidence.",
            "Visual audit complements aggregate metrics.",
        ],
        9.62,
        3.18,
        2.47,
        2.15,
        size=18,
    )
    add_takeaway(slide, "Bounding boxes support review and repair workflows.", y=6.34, size=18)


def slide_failures(prs, assets):
    slide = content_slide(prs, 14, "Representative Failure Cases", "Qualitative audit")
    add_panel_grid(
        slide,
        assets["failure"],
        ["Mouse bite / Short / Spur", "Open circuit / Mouse bite", "Spur", "Short + Spur", "Spur", "Mixed subtle defects"],
    )
    add_card(slide, 9.32, 1.34, 3.12, 4.75, fill=PALE_ORANGE, line=ORANGE)
    add_textbox(slide, "MISS-WEIGHTED AUDIT", 9.62, 1.68, 2.52, 0.3, size=16, bold=True, color=ORANGE)
    add_textbox(slide, "5FN + FP", 9.62, 2.18, 2.52, 0.5, size=30, bold=True, color=NAVY)
    add_bullets(
        slide,
        [
            "False negatives receive the highest review priority.",
            "Small and subtle defects remain difficult.",
            "mAP50-95 exposes loose localization hidden by mAP50.",
        ],
        9.62,
        3.0,
        2.48,
        2.48,
        size=18,
        marker=ORANGE,
    )
    add_takeaway(slide, "False negatives are prioritized because missed defects are costly.", y=6.34, fill=PALE_ORANGE, accent=ORANGE, size=18)


def slide_demo(prs):
    slide = content_slide(prs, 15, "Online Deployment Artifact", "Deployment")
    # Keep the media unobstructed and at the exact source aspect ratio.
    video_w = 8.65
    video_h = video_w / (3420 / 2224)
    video_x = 0.67
    video_y = 1.27
    add_card(slide, video_x - 0.08, video_y - 0.08, video_w + 0.16, video_h + 0.16, fill=NAVY, line=NAVY)
    slide.shapes.add_movie(
        str(VIDEO_PATH),
        Inches(video_x),
        Inches(video_y),
        Inches(video_w),
        Inches(video_h),
        poster_frame_image=str(POSTER_PATH),
        mime_type="video/mp4",
    )
    add_card(slide, 9.62, 1.27, 2.97, 5.64, fill=WHITE)
    add_textbox(slide, "WORKFLOW", 9.92, 1.58, 2.34, 0.25, size=16, bold=True, color=TEAL)
    add_textbox(slide, "Upload  →  Detect\nInspect  →  Download", 9.92, 2.02, 2.34, 0.8, size=20, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_rect(slide, 9.92, 3.02, 2.34, 0.03, LINE)
    add_textbox(slide, "DEPLOYMENT SCOPE", 9.92, 3.33, 2.34, 0.25, size=16, bold=True, color=ORANGE)
    add_textbox(slide, "Packaging and usability", 9.92, 3.78, 2.34, 0.54, size=17, bold=True, color=NAVY)
    add_textbox(slide, "Annotations, class counts, and confidence values", 9.92, 4.44, 2.34, 0.74, size=16, color=INK)
    add_textbox(slide, "Not an embedded-device benchmark", 9.92, 5.42, 2.34, 0.58, size=16, bold=True, color=ORANGE)
    add_textbox(slide, "No Jetson or TensorRT latency claims", 9.92, 6.16, 2.34, 0.54, size=16, bold=True, color=MUTED)


def slide_limitations(prs):
    slide = content_slide(prs, 16, "Limitations and Future Work", "Scope")
    add_card(slide, 0.77, 1.37, 5.75, 4.95, fill=WHITE)
    add_textbox(slide, "CURRENT SCOPE", 1.08, 1.72, 5.1, 0.27, size=18, bold=True, color=ORANGE)
    add_bullets(
        slide,
        [
            "V100 latency, not edge-device latency",
            "Same split, but not fully resource-normalized",
            "In-distribution data, not factory-domain validation",
        ],
        1.1,
        2.3,
        4.95,
        2.75,
        size=21,
        marker=ORANGE,
        gap=14,
    )
    add_card(slide, 6.82, 1.37, 5.75, 4.95, fill=PALE_GREEN, line=TEAL)
    add_textbox(slide, "NEXT STEPS", 7.13, 1.72, 5.1, 0.27, size=18, bold=True, color=GREEN)
    add_bullets(
        slide,
        [
            "Real factory captures and lighting variation",
            "Resource-normalized training and explicit latency budgets",
            "High-resolution crop reinspection",
            "Jetson or TensorRT benchmarking when hardware is available",
        ],
        7.15,
        2.3,
        4.95,
        3.15,
        size=20,
        marker=TEAL,
        gap=10,
    )


def slide_takeaways(prs):
    slide = content_slide(prs, 17, "Key Takeaways", "Close")
    cards = [
        ("PRACTICAL RESULT", "YOLO11s-1280", "Higher recall and mAP50-95\nLower measured latency", TEAL, PALE_GREEN),
        ("CONTROLLED RESULT", "YOLO11s-640", "Competitive at matched resolution\nSlightly stronger recall and strict mAP", BLUE, PALE_BLUE),
        ("DEPLOYMENT LESSON", "Compare carefully", "Resolution and latency matter\nSeparate practical results from ablations", ORANGE, PALE_ORANGE),
    ]
    for index, (label, value, detail, accent, fill) in enumerate(cards):
        x = 0.72 + index * 4.15
        add_card(slide, x, 1.38, 3.72, 2.55, fill=fill, line=accent)
        add_textbox(slide, label, x + 0.25, 1.69, 3.2, 0.25, size=16, bold=True, color=accent, align=PP_ALIGN.CENTER)
        add_textbox(slide, value, x + 0.25, 2.18, 3.2, 0.43, size=27, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_textbox(slide, detail, x + 0.28, 2.91, 3.15, 0.65, size=17, color=INK, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "A deployment-oriented benchmark should separate completed configurations from controlled ablations.", y=4.43, size=19)
    add_textbox(slide, "Thank you. I welcome your questions.", 1.0, 5.55, 11.33, 0.52, size=31, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Aditya Varun Dhayapulay  |  California State University, Fullerton", 1.0, 6.3, 11.33, 0.28, size=16, color=MUTED, align=PP_ALIGN.CENTER)


def copy_speaker_notes(prs):
    source_path = next((path for path in SOURCE_DECKS if path.exists()), None)
    if source_path is None:
        return
    source = Presentation(source_path)
    for index, slide in enumerate(prs.slides):
        if index >= len(source.slides):
            break
        text = source.slides[index].notes_slide.notes_text_frame.text.strip()
        if text:
            slide.notes_slide.notes_text_frame.text = text


def build():
    for required in (VIDEO_PATH, POSTER_PATH):
        if not required.exists():
            raise FileNotFoundError(required)
    assets = make_assets()
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide_title(prs, assets)
    slide_motivation(prs)
    slide_question(prs)
    slide_contributions(prs)
    slide_dataset(prs)
    slide_protocol(prs)
    slide_models(prs)
    slide_fairness(prs)
    slide_practical(prs)
    slide_matched(prs)
    slide_highres(prs)
    slide_per_class(prs)
    slide_successes(prs, assets)
    slide_failures(prs, assets)
    slide_demo(prs)
    slide_limitations(prs)
    slide_takeaways(prs)
    copy_speaker_notes(prs)

    # Keep Q&A backup material in the same deck for instant access.
    try:
        from tools.build_escs26_qa_backup_slides import add_backup_slides
    except ModuleNotFoundError:
        from build_escs26_qa_backup_slides import add_backup_slides

    add_backup_slides(prs)
    prs.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
